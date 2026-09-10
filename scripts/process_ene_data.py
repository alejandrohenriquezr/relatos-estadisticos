#!/usr/bin/env python3
"""Convierte libros oficiales de la ENE en el JSON liviano usado por la web.

El script recibe un directorio de insumos en vez de depender de una carpeta de
una sesión de trabajo. Las observaciones que no están contenidas en esos libros
pueden agregarse mediante un JSON suplementario explícito; no se incorporan
cifras manuales ocultas en el código.

Archivos esperados en el directorio de entrada:
- indicadores_principales.xlsx (también acepta "indicadores_principales (1).xlsx")
- rama.xlsx
- categoria.xlsx
- ajuste_estacional_historico.xlsx

Uso:
    python scripts/process_ene_data.py datos/ene
    python scripts/process_ene_data.py datos/ene --supplement datos/ene/supplement.json
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


DEFAULT_OUTPUT = Path(__file__).resolve().parents[1] / "public" / "ene-data.json"


def clean_number(value: Any) -> float | None:
    """Devuelve números redondeados y descarta códigos de calidad no numéricos."""
    return round(float(value), 4) if isinstance(value, (int, float)) else None


def resolve_file(input_dir: Path, *names: str) -> Path:
    """Resuelve uno de los nombres admitidos dentro del directorio portable."""
    for name in names:
        candidate = input_dir / name
        if candidate.is_file():
            return candidate
    expected = ", ".join(names)
    raise FileNotFoundError(f"Falta un insumo ENE en {input_dir}: {expected}")


def principal_series(input_dir: Path) -> dict[str, list[dict[str, Any]]]:
    """Extrae las series nacionales para total, mujeres y hombres."""
    source = resolve_file(
        input_dir,
        "indicadores_principales.xlsx",
        "indicadores_principales (1).xlsx",
    )
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        result: dict[str, list[dict[str, Any]]] = {}
        for sheet, label in (("AS", "Total"), ("M", "Mujeres"), ("H", "Hombres")):
            ws = workbook[sheet]
            rows: list[dict[str, Any]] = []
            for row in ws.iter_rows(min_row=8, values_only=True):
                year, quarter = row[0], row[1]
                if not isinstance(year, int) or not isinstance(quarter, str):
                    continue
                pet, labor, employed, unemployed = map(
                    clean_number, (row[3], row[5], row[7], row[9])
                )
                ceased, first_job = map(clean_number, (row[11], row[13]))
                unemployment_rate, employment_rate, participation = map(
                    clean_number, (row[23], row[25], row[27])
                )
                if None in (
                    pet,
                    labor,
                    employed,
                    unemployed,
                    unemployment_rate,
                    employment_rate,
                    participation,
                ):
                    continue
                rows.append(
                    {
                        "year": year,
                        "quarter": " ".join(quarter.split()),
                        "pet": pet,
                        "labor": labor,
                        "employed": employed,
                        "unemployed": unemployed,
                        "ceased": ceased,
                        "firstJob": first_job,
                        "participation": participation,
                        "employmentRate": employment_rate,
                        "unemploymentRate": unemployment_rate,
                    }
                )
            result[label] = rows
        return result
    finally:
        workbook.close()


def latest_breakdown(input_dir: Path, filename: str, sheet: str) -> dict[str, Any]:
    """Extrae la última distribución disponible de ramas o categorías."""
    source = resolve_file(input_dir, filename)
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        ws = workbook[sheet]
        header_row = 4 if filename in {"rama.xlsx", "categoria.xlsx"} else 6
        headers = [ws.cell(header_row, col).value for col in range(1, ws.max_column + 1)]
        latest = None
        for row in ws.iter_rows(min_row=header_row + 2, values_only=True):
            if isinstance(row[0], int) and isinstance(row[1], str):
                latest = row
        if latest is None:
            return {"period": "", "items": []}
        items = []
        for index in range(2, len(latest)):
            value = clean_number(latest[index])
            header = headers[index] if index < len(headers) else None
            if value is not None and isinstance(header, str) and header.strip():
                items.append({"label": " ".join(header.split()), "value": value})
        return {"period": f"{latest[1]} {latest[0]}", "items": items[:18]}
    finally:
        workbook.close()


def seasonal_unemployment(input_dir: Path) -> list[dict[str, Any]]:
    """Obtiene la tasa desestacionalizada del libro histórico de ajuste."""
    source = resolve_file(input_dir, "ajuste_estacional_historico.xlsx")
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        ws = workbook["tasa_as"]
        records = []
        for row in ws.iter_rows(min_row=8, values_only=True):
            year, quarter = row[0], row[1]
            adjusted = clean_number(row[155]) if len(row) >= 156 else None
            if isinstance(year, int) and isinstance(quarter, str) and adjusted is not None:
                records.append(
                    {
                        "year": year,
                        "quarter": " ".join(quarter.split()),
                        "value": adjusted,
                    }
                )
        return list(reversed(records))
    finally:
        workbook.close()


def contribution_rows(
    input_dir: Path,
    filename: str,
    excluded: set[str] | None = None,
    display_names: dict[str, str] | None = None,
) -> list[dict[str, Any]]:
    """Calcula incidencias positivas anuales desde ramas o categorías."""
    source = resolve_file(input_dir, filename)
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        ws = workbook["AS"]
        excluded = excluded or set()
        display_names = display_names or {}
        names = []
        for col in range(5, ws.max_column + 1, 2):
            header = ws.cell(6, col).value
            if not isinstance(header, str):
                continue
            normalized = " ".join(header.split())
            if normalized in excluded:
                continue
            names.append((col + 1, display_names.get(normalized, normalized)))

        rows: dict[tuple[int, str], tuple[Any, ...]] = {}
        for row in ws.iter_rows(min_row=8, values_only=True):
            if isinstance(row[0], int) and isinstance(row[1], str):
                rows[(row[0], " ".join(row[1].split()))] = row

        output = []
        for (year, quarter), row in rows.items():
            previous = rows.get((year - 1, quarter))
            if previous is None or not isinstance(previous[3], (int, float)):
                continue
            items = []
            for excel_col, label in names:
                index = excel_col - 1
                current_value = row[index] if index < len(row) else None
                previous_value = previous[index] if index < len(previous) else None
                if not isinstance(current_value, (int, float)) or not isinstance(
                    previous_value, (int, float)
                ) or not previous_value:
                    continue
                change = (current_value / previous_value - 1) * 100
                incidence = (current_value - previous_value) / previous[3] * 100
                if incidence > 0:
                    items.append(
                        {
                            "label": label,
                            "change": round(change, 2),
                            "incidence": round(incidence, 3),
                        }
                    )
            items.sort(key=lambda item: item["incidence"], reverse=True)
            output.append({"year": year, "quarter": quarter, "items": items[:3]})
        return output
    finally:
        workbook.close()


def sector_contributions(input_dir: Path) -> list[dict[str, Any]]:
    """Calcula las ramas con mayor incidencia positiva anual en ocupación."""
    return contribution_rows(input_dir, "rama.xlsx")


def category_contributions(input_dir: Path) -> list[dict[str, Any]]:
    """Calcula categorías ocupacionales específicas con incidencia positiva."""
    excluded = {
        "Independientes (Total) [2]",
        "Dependientes (Total) [3]",
        "Asalariados/as (Total) [4]",
    }
    display_names = {
        "Independientes (Empleadores/as)": "Personas empleadoras",
        "Independientes (Trabajadores/as por cuenta propia)": "Trabajadores/as por cuenta propia",
        "Independientes (Familiares no remunerados)": "Familiares no remunerados",
        "Asalariados/as (Sector privado)": "Personas asalariadas del sector privado",
        "Asalariados/as (Sector público) [5]": "Personas asalariadas del sector público",
        "Personal de servicio doméstico (Total) [6]": "Personal de servicio doméstico",
    }
    return contribution_rows(input_dir, "categoria.xlsx", excluded, display_names)


def load_supplement(path: Path | None) -> dict[str, Any]:
    """Carga observaciones externas sólo cuando se entregan explícitamente."""
    if path is None:
        return {}
    if not path.is_file():
        raise FileNotFoundError(f"No existe el suplemento ENE: {path}")
    payload = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(payload, dict):
        raise ValueError("El suplemento ENE debe ser un objeto JSON")
    return payload


def main() -> None:
    """Construye el snapshot a partir de insumos declarados y verificables."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "input_dir",
        type=Path,
        help="Directorio que contiene los libros oficiales requeridos",
    )
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument(
        "--supplement",
        type=Path,
        help="JSON opcional para observaciones no presentes en los libros",
    )
    args = parser.parse_args()
    input_dir = args.input_dir.expanduser().resolve()
    if not input_dir.is_dir():
        raise FileNotFoundError(f"No existe el directorio de insumos ENE: {input_dir}")

    principal = principal_series(input_dir)
    if not principal.get("Total"):
        raise ValueError("La serie principal ENE quedó vacía")
    supplement = load_supplement(args.supplement)
    latest = principal["Total"][-1]

    sectors = sector_contributions(input_dir)
    sectors.extend(supplement.get("sectorContributions", []))
    categories = category_contributions(input_dir)
    categories.extend(supplement.get("categoryContributions", []))

    payload = {
        "series": principal,
        "seasonal": seasonal_unemployment(input_dir),
        "sectorContributions": sectors,
        "categoryContributions": categories,
        "absentEmployment": supplement.get("absentEmployment", []),
        "branches": latest_breakdown(input_dir, "rama.xlsx", "AS"),
        "categories": latest_breakdown(input_dir, "categoria.xlsx", "AS"),
        "metadata": {
            "source": "Instituto Nacional de Estadísticas de Chile, Encuesta Nacional de Empleo",
            "updated": f"{latest['quarter']} {latest['year']}",
            "supplement": str(args.supplement) if args.supplement else None,
            "qualityNotes": {
                "a": "Estimación poco fiable; debe utilizarse con precaución.",
                "b": "Estimación no fiable.",
            },
        },
    }

    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(payload, ensure_ascii=False, separators=(",", ":")),
        encoding="utf-8",
    )
    print(
        json.dumps(
            {
                "output": str(args.output),
                "latest": payload["metadata"]["updated"],
                "principalObservations": len(principal["Total"]),
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
