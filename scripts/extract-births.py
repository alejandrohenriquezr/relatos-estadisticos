#!/usr/bin/env python3
"""Convierte la hoja Nacimientos del libro oficial de estadísticas vitales a JSON.

Acepta una ruta local o una URL. Por defecto usa la serie histórica oficial
1992-2025(p), cuya compatibilidad con la estructura del sitio fue verificada
contra el parser productivo.
"""

from __future__ import annotations

import argparse
import json
import re
import tempfile
import urllib.request
from contextlib import contextmanager
from pathlib import Path
from typing import Iterator

from openpyxl import load_workbook


DEFAULT_SOURCE = (
    "https://www.ine.gob.cl/docs/default-source/nacimientos-matrimonios-y-defunciones/"
    "cuadros-estadisticos/series-hist%C3%B3ricas/series-vitales-1992-2025(p).xlsx"
    "?sfvrsn=bfbe614_4"
)
DEFAULT_OUTPUT = Path(__file__).resolve().parents[1] / "public" / "births-data.json"


@contextmanager
def source_path(source: str) -> Iterator[Path]:
    """Materializa una URL XLSX o entrega una ruta local validada."""
    if not source.lower().startswith(("http://", "https://")):
        path = Path(source).expanduser().resolve()
        if not path.is_file():
            raise FileNotFoundError(f"No existe el libro de entrada: {path}")
        yield path
        return

    request = urllib.request.Request(
        source,
        headers={"User-Agent": "INE-Relatos-source-extractor/1.0"},
    )
    temporary = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    path = Path(temporary.name)
    try:
        with urllib.request.urlopen(request, timeout=120) as response:
            while chunk := response.read(1024 * 1024):
                temporary.write(chunk)
        temporary.close()
        if path.read_bytes()[:4] != b"PK\x03\x04":
            raise ValueError("La fuente de nacimientos no devolvió un XLSX válido")
        yield path
    finally:
        temporary.close()
        path.unlink(missing_ok=True)


def build_payload(source: Path, source_label: str) -> dict:
    """Extrae las observaciones anuales y los grupos de edad disponibles."""
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        if "Nacimientos" not in workbook.sheetnames:
            raise ValueError("El libro no contiene la hoja 'Nacimientos'")
        worksheet = workbook["Nacimientos"]
        headers = [
            cell.value
            for cell in next(worksheet.iter_rows(min_row=1, max_row=1))
        ]
        age_columns = [
            (
                index,
                str(label)
                .replace("Nacimientos de mujeres de ", "")
                .replace("Nacimientos de mujeres ", ""),
            )
            for index, label in enumerate(headers)
            if label
            and str(label).startswith("Nacimientos de mujeres")
            and "no especificada" not in str(label)
        ]
        series = []
        for row in worksheet.iter_rows(min_row=2, values_only=True):
            if row[0] is None:
                break
            match = re.search(r"\d{4}", str(row[0]))
            if not match:
                continue
            series.append(
                {
                    "year": int(match.group()),
                    "provisional": "(p)" in str(row[0]),
                    "observed": int(row[1]),
                    "men": int(row[3]),
                    "women": int(row[4]),
                    "masculinity": float(row[6]),
                    "ages": [
                        {"label": label, "value": int(row[index] or 0)}
                        for index, label in age_columns
                    ],
                }
            )
    finally:
        workbook.close()

    if not series:
        raise ValueError("No se encontraron nacimientos válidos")
    if series[-1]["year"] < 2025:
        raise ValueError("La fuente no contiene la serie provisional 2025 esperada")
    return {
        "source": (
            Path(source_label).name
            if not source_label.startswith("http")
            else source_label.rsplit("/", 1)[-1].split("?", 1)[0]
        ),
        "sourceUrl": (
            source_label
            if source_label.startswith(("http://", "https://"))
            else None
        ),
        "series": series,
    }


def main() -> None:
    """Transforma el origen elegido y escribe el snapshot utilizado por la web."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", nargs="?", default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    with source_path(args.source) as local_source:
        payload = build_payload(local_source, args.source)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(payload, ensure_ascii=False, separators=(",", ":")),
        encoding="utf-8",
    )
    print(
        json.dumps(
            {
                "output": str(args.output),
                "years": len(payload["series"]),
                "lastYear": payload["series"][-1]["year"],
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
