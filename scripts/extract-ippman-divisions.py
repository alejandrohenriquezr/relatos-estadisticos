#!/usr/bin/env python3
"""Extrae divisiones manufactureras desde el cuadro oficial del IPP.

El origen puede ser una ruta local o una URL directa. Si se omite, se utiliza
la fuente estable publicada por el INE para industria manufacturera, base anual
2019=100.
"""

from __future__ import annotations

import argparse
import json
import tempfile
import urllib.request
from contextlib import contextmanager
from pathlib import Path
from typing import Iterator

from openpyxl import load_workbook


DEFAULT_SOURCE = (
    "https://www.ine.gob.cl/docs/default-source/indice-de-precios-de-productor/"
    "cuadros-estadisticos/base-anual-2019-100/industria-manufacturera-xlsx.xlsx"
)
DEFAULT_OUTPUT = (
    Path(__file__).resolve().parents[1] / "public" / "ippman-divisions.json"
)


@contextmanager
def source_path(source: str) -> Iterator[Path]:
    """Materializa una URL XLSX o valida una ruta local."""
    if not source.lower().startswith(("http://", "https://")):
        path = Path(source).expanduser().resolve()
        if not path.is_file():
            raise FileNotFoundError(f"No existe el libro de entrada: {path}")
        yield path
        return

    request = urllib.request.Request(
        source,
        headers={"User-Agent": "INE-Relatos-source-extractor/1.0"},
    )
    temporary = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    path = Path(temporary.name)
    try:
        with urllib.request.urlopen(request, timeout=120) as response:
            while chunk := response.read(1024 * 1024):
                temporary.write(chunk)
        temporary.close()
        if path.read_bytes()[:4] != b"PK\x03\x04":
            raise ValueError("La fuente IPP manufactura no devolvió un XLSX válido")
        yield path
    finally:
        temporary.close()
        path.unlink(missing_ok=True)


def build_rows(source: Path) -> list[dict]:
    """Conserva sólo el nivel división de la hoja IPP_Manufactura."""
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        if "IPP_Manufactura" not in workbook.sheetnames:
            raise ValueError("El libro no contiene la hoja 'IPP_Manufactura'")
        worksheet = workbook["IPP_Manufactura"]
        rows: list[dict] = []

        for row in worksheet.iter_rows(min_row=6, values_only=True):
            (
                year,
                month,
                division,
                group,
                _class,
                subclass,
                product,
                label,
                _,
                index,
                monthly,
                accumulated,
                annual,
            ) = row[:13]
            if division is None or group is not None:
                continue
            if not isinstance(year, (int, float)) or not isinstance(month, (int, float)):
                continue
            rows.append(
                {
                    "year": int(year),
                    "month": int(month),
                    "division": int(division),
                    "label": str(label),
                    "index": float(index),
                    "monthly": float(monthly),
                    "accumulated": float(accumulated),
                    "annual": float(annual),
                }
            )
    finally:
        workbook.close()

    if not rows:
        raise ValueError("No se encontraron divisiones manufactureras válidas")
    return rows


def main() -> None:
    """Transforma el libro seleccionado y escribe la instantánea JSON."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", nargs="?", default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    with source_path(args.source) as local_source:
        rows = build_rows(local_source)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(rows, ensure_ascii=False, separators=(",", ":")),
        encoding="utf-8",
    )
    latest = max((row["year"], row["month"]) for row in rows)
    print(
        json.dumps(
            {
                "output": str(args.output),
                "observations": len(rows),
                "lastPeriod": f"{latest[0]:04d}-{latest[1]:02d}",
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
