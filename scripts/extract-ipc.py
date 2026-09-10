#!/usr/bin/env python3
"""Extrae IPC general y divisiones desde el Excel oficial del INE.

Acepta una ruta local o una URL. Si no se indica origen, usa exactamente la
misma fuente estable que `app/api/ipc-data/route.ts`: IPC base anual 2023=100,
series de tiempo. El período de actualización se calcula desde la última
observación del archivo y no queda fijado manualmente en el código.
"""

from __future__ import annotations

import argparse
import json
import tempfile
import urllib.request
from contextlib import contextmanager
from pathlib import Path
from typing import Iterator

import openpyxl


DEFAULT_SOURCE = (
    "https://www.ine.gob.cl/docs/default-source/%C3%ADndice-de-precios-al-consumidor/"
    "cuadros-estadisticos/base-anual-2023_100/series-de-tiempo/ipc-xls.xlsx"
)
DEFAULT_OUTPUT = Path(__file__).resolve().parents[1] / "public" / "ipc-data.json"
MONTHS = {
    1: "Enero",
    2: "Febrero",
    3: "Marzo",
    4: "Abril",
    5: "Mayo",
    6: "Junio",
    7: "Julio",
    8: "Agosto",
    9: "Septiembre",
    10: "Octubre",
    11: "Noviembre",
    12: "Diciembre",
}


@contextmanager
def source_path(source: str) -> Iterator[Path]:
    """Materializa una URL en un archivo temporal o valida una ruta local."""
    if not source.lower().startswith(("http://", "https://")):
        path = Path(source).expanduser().resolve()
        if not path.is_file():
            raise FileNotFoundError(f"No existe el libro de entrada: {path}")
        yield path
        return

    request = urllib.request.Request(
        source,
        headers={"User-Agent": "INE-Relatos-source-extractor/1.0"},
    )
    temporary = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    path = Path(temporary.name)
    try:
        with urllib.request.urlopen(request, timeout=120) as response:
            while chunk := response.read(1024 * 1024):
                temporary.write(chunk)
        temporary.close()
        if path.read_bytes()[:4] != b"PK\x03\x04":
            raise ValueError("La fuente IPC no devolvió un XLSX válido")
        yield path
    finally:
        temporary.close()
        path.unlink(missing_ok=True)


def build_payload(source: Path, source_label: str) -> dict:
    """Transforma el libro al contrato JSON usado por la interfaz."""
    workbook = openpyxl.load_workbook(source, read_only=True, data_only=True)
    try:
        if "IPC 2023=100" not in workbook.sheetnames:
            raise ValueError("El libro no contiene la hoja 'IPC 2023=100'")
        sheet = workbook["IPC 2023=100"]
        series: list[dict] = []

        for row in sheet.iter_rows(min_row=5, values_only=True):
            year, month, division, group, class_, subclass, product, label = row[:8]
            is_general = label == "IPC General"
            is_division = division is not None and all(
                value is None for value in (group, class_, subclass, product)
            )
            if not (is_general or is_division):
                continue
            if not isinstance(year, (int, float)) or not isinstance(month, (int, float)):
                continue
            series.append(
                {
                    "year": int(year),
                    "month": int(month),
                    "division": 0 if is_general else int(division),
                    "label": str(label),
                    "weight": row[8],
                    "index": row[9],
                    "monthly": row[10],
                    "accumulated": row[11],
                    "annual": row[12],
                    "monthlyIncidence": row[13],
                }
            )
    finally:
        workbook.close()

    if not series:
        raise ValueError("No se encontraron observaciones IPC válidas")
    latest_year, latest_month = max((item["year"], item["month"]) for item in series)
    updated = f"{MONTHS[latest_month]} de {latest_year}"
    return {
        "base": "2023=100",
        "updated": updated,
        "source": source_label,
        "series": series,
    }


def main() -> None:
    """Resuelve el origen, transforma y escribe un JSON UTF-8 reproducible."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "source",
        nargs="?",
        default=DEFAULT_SOURCE,
        help="Ruta local o URL del XLSX; por defecto se usa la fuente oficial",
    )
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    with source_path(args.source) as local_source:
        payload = build_payload(local_source, args.source)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(payload, ensure_ascii=False, separators=(",", ":")),
        encoding="utf-8",
    )
    print(
        json.dumps(
            {
                "output": str(args.output),
                "updated": payload["updated"],
                "observations": len(payload["series"]),
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
