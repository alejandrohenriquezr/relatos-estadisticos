#!/usr/bin/env python3
"""Extrae las series policiales desde un libro oficial del INE.

El origen puede ser una ruta local o una URL directa a un XLSX. No se fija un
año máximo: el año publicado se obtiene del contenido del libro, por lo que el
mismo extractor sirve para futuras actualizaciones mientras se conserve la
estructura de hojas utilizada por el INE.

Uso:
    python scripts/extract-police.py archivo.xlsx
    python scripts/extract-police.py https://.../archivo.xlsx
    python scripts/extract-police.py archivo.xlsx --output salida.json
"""

from __future__ import annotations

import argparse
import json
import re
import tempfile
import urllib.request
from contextlib import contextmanager
from pathlib import Path
from typing import Iterator

import openpyxl


DEFAULT_OUTPUT = Path(__file__).resolve().parents[1] / "public" / "police-data.json"


def clean_region(value: object) -> str:
    """Normaliza glosas territoriales y elimina marcas de nota al pie."""
    label = str(value or "").replace("\u00a0", " ").strip()
    label = re.sub(r"/\d+$", "", label).strip()
    return label.upper()


def numeric(value: object) -> int | float | None:
    """Conserva sólo celdas numéricas del libro."""
    return value if isinstance(value, (int, float)) else None


@contextmanager
def source_path(source: str) -> Iterator[Path]:
    """Entrega una ruta local para un archivo local o una URL remota."""
    if not source.lower().startswith(("http://", "https://")):
        path = Path(source).expanduser().resolve()
        if not path.is_file():
            raise FileNotFoundError(f"No existe el libro de entrada: {path}")
        yield path
        return

    request = urllib.request.Request(
        source,
        headers={"User-Agent": "INE-Relatos-source-extractor/1.0"},
    )
    temporary = tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False)
    path = Path(temporary.name)
    try:
        with urllib.request.urlopen(request, timeout=120) as response:
            while chunk := response.read(1024 * 1024):
                temporary.write(chunk)
        temporary.close()
        if path.read_bytes()[:4] != b"PK\x03\x04":
            raise ValueError("La URL no devolvió un archivo XLSX válido")
        yield path
    finally:
        temporary.close()
        path.unlink(missing_ok=True)


def extract_sheet(workbook, sheet_name: str, combined: bool = False) -> list[dict]:
    """Extrae una serie anual total y regional desde una hoja histórica."""
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"El libro no contiene la hoja requerida: {sheet_name}")

    sheet = workbook[sheet_name]
    # Fila 4: Año/Caso, Total, Variación y luego las regiones. La columna de
    # variación se excluye explícitamente para que nunca aparezca como región.
    region_headers = [clean_region(cell.value) for cell in sheet[4][3:]]
    records: list[dict] = []

    for row in sheet.iter_rows(min_row=5, values_only=True):
        label = row[0]
        if combined and (
            not isinstance(label, str) or not label.strip().startswith("Denuncias ")
        ):
            continue
        match = re.search(r"(20\d{2})", str(label))
        if not match:
            continue
        year = int(match.group(1))
        regional = {
            region: numeric(value)
            for region, value in zip(region_headers, row[3:])
            if region and region != "NONE" and not region.startswith("VARIACIÓN")
        }
        records.append(
            {
                "year": year,
                "total": numeric(row[1]),
                "regions": regional,
            }
        )
    return records


def build_payload(workbook) -> dict:
    """Construye el contrato JSON consumido por la página policial."""
    institutions = {
        "carabineros": {
            "label": "Carabineros de Chile",
            "series": {
                "denuncias": extract_sheet(workbook, "1", combined=True),
                "detenidos": extract_sheet(workbook, "7"),
                "victimas": extract_sheet(workbook, "17"),
            },
        },
        "pdi": {
            "label": "Policía de Investigaciones de Chile",
            "series": {
                "denuncias": extract_sheet(workbook, "23"),
                "detenidos": extract_sheet(workbook, "28"),
                "victimas": extract_sheet(workbook, "37"),
            },
        },
    }
    years = [
        record["year"]
        for institution in institutions.values()
        for series in institution["series"].values()
        for record in series
    ]
    if not years:
        raise ValueError("No se encontraron observaciones anuales en el libro")
    updated = max(years)
    for institution in institutions.values():
        for series in institution["series"].values():
            if not any(record["year"] == updated for record in series):
                raise ValueError(
                    f"La serie policial no contiene el último año esperado: {updated}"
                )
    return {"updated": updated, "institutions": institutions}


def main() -> None:
    """Lee el libro, valida las seis series y escribe el snapshot."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", help="Ruta local o URL directa al XLSX oficial")
    parser.add_argument(
        "--output",
        type=Path,
        default=DEFAULT_OUTPUT,
        help=f"JSON de salida (predeterminado: {DEFAULT_OUTPUT})",
    )
    args = parser.parse_args()

    with source_path(args.source) as source:
        workbook = openpyxl.load_workbook(source, data_only=True, read_only=True)
        try:
            payload = build_payload(workbook)
        finally:
            workbook.close()

    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(
        json.dumps(
            {"output": str(args.output), "updated": payload["updated"]},
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
